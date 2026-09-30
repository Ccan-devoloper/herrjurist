"""Schreibt THEMENPLAN-780.md, themenplan-780.csv, themenreserve.csv und die Excel-Fassung aus plan.json/reserve.json.

    python3 ausgabe.py ../THEMENPLAN-780.md . LexVerse_Themenplan_780.xlsx
"""
import json, csv, collections, statistics, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ZIEL_MD, ZIEL_DIR, ZIEL_XLSX = sys.argv[1], sys.argv[2], sys.argv[3]
P = json.load(open("plan.json")); R = json.load(open("reserve.json")); ST = json.load(open("stats.json"))
SP = ["nr", "jahr", "woche", "woche_im_jahr", "tag", "reihe", "gebiet", "teilgebiet", "titel", "hook", "kernfrage", "normen",
      "leitentscheidung", "examen", "format", "relevanz", "breite", "klassiker", "score", "voraussetzung", "fundstelle", "rechtsstand"]
KOPF = ["Nr.", "Jahr", "Woche", "Woche im Jahr", "Tag", "Reihe", "Gebiet", "Teilgebiet", "Titel", "Einstieg (Hook)", "Kernfrage",
        "Normen", "Leitentscheidung", "Examen", "Format", "Relevanz (1–5)", "Breite (1–5)", "Klassiker", "Punkte",
        "Voraussetzung", "Fundstelle im Repetitorium", "Rechtsstand/Länder"]


def csv_schreiben(pfad, zeilen, spalten):
    with open(pfad, "w", newline="", encoding="utf-8-sig") as f:        # BOM: Excel erkennt Umlaute
        w = csv.writer(f, delimiter=";")
        w.writerow([KOPF[SP.index(s)] for s in spalten])
        for x in zeilen:
            w.writerow(["ja" if x.get(s) is True else "" if x.get(s) is False else x.get(s, "") for s in spalten])


csv_schreiben(f"{ZIEL_DIR}/themenplan-780.csv", P, SP)
RSP = [s for s in SP if s not in ("nr", "jahr", "woche", "woche_im_jahr", "tag", "reihe")]
csv_schreiben(f"{ZIEL_DIR}/themenreserve.csv", R, RSP)

# --- Excel ------------------------------------------------------------------------------------------------------------------
wb = Workbook()
FARBE = {"Zivilrecht": "DCEBFB", "Öffentliches Recht": "E3F4E1", "Strafrecht": "FBE3E0", "2. Examen": "EFE6FA", "Methodik": "FFF4D6"}


def blatt(ws, zeilen, spalten, breiten):
    ws.append([KOPF[SP.index(s)] for s in spalten])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F2A44")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for x in zeilen:
        ws.append(["ja" if x.get(s) is True else "" if x.get(s) is False else x.get(s, "") for s in spalten])
        f = PatternFill("solid", fgColor=FARBE[x["gebiet"]])
        for c in ws[ws.max_row]:
            c.fill = f; c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, b in enumerate(breiten, 1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions


ws = wb.active; ws.title = "Plan 780"
BR = {"nr": 6, "jahr": 6, "woche": 7, "woche_im_jahr": 8, "tag": 5, "reihe": 14, "gebiet": 16, "teilgebiet": 18, "titel": 55,
      "hook": 55, "kernfrage": 50, "normen": 22, "leitentscheidung": 26, "examen": 7, "format": 14, "relevanz": 9, "breite": 9,
      "klassiker": 9, "score": 7, "voraussetzung": 40, "fundstelle": 35, "rechtsstand": 40}
blatt(ws, P, SP, [BR[s] for s in SP])
for j in range(1, 6):
    blatt(wb.create_sheet(f"Jahr {j}"), [x for x in P if x["jahr"] == j], SP, [BR[s] for s in SP])
blatt(wb.create_sheet("Reserve"), R, RSP, [BR[s] for s in RSP])
wb.save(ZIEL_XLSX)

# --- Markdown -------------------------------------------------------------------------------------------------------------
def md(s):
    return str(s).replace("|", "/").replace("\n", " ")


c = collections.Counter(x["gebiet"] for x in P)
GEB = ["Zivilrecht", "Strafrecht", "Öffentliches Recht", "2. Examen", "Methodik"]
jahrstat = []
for j in range(1, 6):
    l = [x for x in P if x["jahr"] == j]
    jahrstat.append((j, round(statistics.mean(x["score"] for x in l), 1), sum(x["klassiker"] for x in l),
                     sum(x["breite"] >= 4 for x in l)))
rs = [x for x in P if x["rechtsstand"] and x["rechtsstand"] != "stabil" and not x["rechtsstand"].startswith("Länder")]
laender = [x for x in P if x["rechtsstand"].startswith("Länder")]
tg = collections.Counter((x["gebiet"], x["teilgebiet"]) for x in P)

L = []
a = L.append
a("# LexVerse · Themenplan für 5 Jahre (780 Folgen)\n")
a("**Stand 30.09.2026.** Dieser Plan ersetzt den bisherigen Jahresplan [SERIENPLAN-156.md](SERIENPLAN-156.md) vollständig; auch die bereits produzierten Folgen werden neu gemacht. "
  "Grundlage ist das vollständige Vollrepetitorium des Kanalinhabers (`Jura_Gesamtwissen…UPDATE30…UNCERTIFIED…html`, 15,4 Mio. Zeichen, 13.356 Überschriften), "
  "ausgewertet nach Rechtsgebieten, dazu eine Web-Recherche zur YouTube-Nachfrage ([themenplanung/recherche_youtube.md](themenplanung/recherche_youtube.md)). "
  "Vollständige Tabelle mit Einstieg, Kernfrage, Normen und Fundstelle: [themenplanung/themenplan-780.csv](themenplanung/themenplan-780.csv); " + str(len(R)) + " bewertete Ersatzthemen: [themenplanung/themenreserve.csv](themenplanung/themenreserve.csv). "
  "Neu erzeugen: `cd youtube/themenplanung && python3 plan_bauen.py && python3 ausgabe.py ../THEMENPLAN-780.md . LexVerse_Themenplan_780.xlsx` (benötigt `openpyxl`).\n")
a("## Kurzfassung\n")
a("- **3 Folgen pro Woche, 52 Wochen, 5 Jahre = 780 Folgen** à 5–7 Minuten (Zielwert 5–6 Minuten Inhalt; Lernvideos verlieren laut edX-Studie ab etwa 6 Minuten deutlich an Aufmerksamkeit).")
a("- **Feste Wochenstruktur** mit drei wiedererkennbaren Reihen:")
a("  - **Montag · Der Fall** – berühmte Leitentscheidung oder Alltagsfall als Aufhänger (Raser-Fälle, Haustyrann, Brokdorf, Trierer Weinversteigerung). Das ist das Zugpferd für Klicks und neue Abonnenten.")
a("  - **Mittwoch · Examenswissen** – Prüfungsschema, Streitstand oder Abgrenzung. Diese Videos werden das ganze Jahr über gesucht (Notwehr-Schema, Anfechtungsklage, § 437 BGB).")
a("  - **Freitag · Klausurpraxis** – im Wechsel 2. Examen (etwa jede zweite bis dritte Woche), Methodik (7× pro Jahr, zu Semesterbeginn und in Klausurphasen) und Abgrenzungen/Klausurfehler aus dem 1. Examen.")
a("- **Jede Woche mindestens zwei verschiedene Rechtsgebiete**; Zivil-, Straf- und Öffentliches Recht wechseln sich ab.")
a("- **Jahr 1 bringt die stärksten Themen und den Grundstock**: 44 Einsteigervideos (Gutachtenstil, Straftataufbau, Notwehr, Diebstahl, Betrug, Grundrechtsprüfung, Verfassungsbeschwerde, Anfechtungsklage, Vertragsschluss, Anfechtung, § 437, § 823) in Lernreihenfolge, eingebettet zwischen die zugkräftigsten Fallvideos.")
a("- **Woche 1 = erste Oktoberwoche** (Semesterstart). Methodik liegt in den Wochen 1–3 (Oktober), 15 (Januar, Klausurphase), 27 (April), 39 und 41 (Juni/Juli). Wer zu einem anderen Zeitpunkt startet, verschiebt nur die Methodik-Wochen entsprechend.\n")
a("## Verteilung\n")
a("| Gebiet | Folgen | Anteil | pro Jahr |")
a("| --- | ---: | ---: | ---: |")
for g in GEB:
    a(f"| {g} | {c[g]} | {c[g] / 7.8:.0f} % | {c[g] // 5} |")
a("| **Summe** | **780** | 100 % | 156 |\n")
a("Strafrecht ist bewusst etwas stärker gewichtet: Laut Jurafuchs-Auswertung des Google-Suchinteresses (Mai 2025–April 2026) liegt Strafrecht mit 272 Suchanfragen je 100.000 Nutzer klar vorn. Zugleich schlägt es die Brücke zwischen Studierenden und Laien (Notwehr, Polizei, Betrug im Netz).\n")
a("<details><summary>Teilgebiete</summary>\n")
a("| Gebiet | Teilgebiet | Folgen |")
a("| --- | --- | ---: |")
for (g, t), n in sorted(tg.items(), key=lambda k: (GEB.index(k[0][0]), -k[1])):
    a(f"| {g} | {t} | {n} |")
a("\n</details>\n")
a("## So wurden die Themen ausgewählt\n")
a("1. **Vollständige Auswertung des Repetitoriums nach Rechtsgebieten.** Ausgewertet wurden die Bände 1–4 (materielles Recht und 2. Examen), Band 8 (Einzelprobleme und Streitstände), Band 9 (Standardfälle), die konsolidierten Bände 11–14, die Problematlanten 17–20 sowie die Klassiker- und Pflichtfalllisten (Band 7 § 9, Band 9 Nr. 75–78). "
  "Die sehr umfangreichen Bände 42–73 und der UPDATE30-Anhang (zusammen gut 12 Mio. Zeichen) sind überwiegend Lexikon-, Normen- und Seitenabgleiche zur Lernplattform. Sie wurden über ihre Gliederung und über Normhäufigkeiten im Gesamttext ausgewertet, nicht Satz für Satz, etwa § 80 VwGO 781×, § 28 VwVfG 506×, § 355 BGB 406× und § 338 StPO 394×.")
a("2. **Rund 1.000 Kandidaten** (Zivilrecht 318, Öffentliches Recht 250, Strafrecht 263, 2. Examen 134, Methodik 36). Jeder Kandidat ist genau ein Problem, das sich in 5–7 Minuten erklären lässt, mit eigenem Alltags-Einstieg. Sachverhalte der Lernplattform wurden nicht übernommen, echte Entscheidungsnamen schon. 11 Dubletten über Gebietsgrenzen hinweg wurden gestrichen.")
a("3. **Bewertung:**")
a("   - Relevanz im Examen (1–5) und Breite, also Interesse über Jurastudierende hinaus (1–5).")
a("   - `Punkte = 2 × Relevanz + 1,5 × Breite + 2 (Klassiker) + 1,5 (Leitfall) bzw. 1 (Alltagsfall) + 0,5 (Strafrecht) + 3 (Kernschema mit Relevanz 5)`.")
a("   - Innerhalb jedes Gebiets bleiben alle Teilgebiete angemessen vertreten.")
a("4. **Feste Bestandteile:**")
a("   - 44 Grundstock-Themen sind verbindlich in Jahr 1 und stehen dort in didaktischer Reihenfolge.")
a("   - 7 Pflichtthemen (etwa Kausalität, fehlgeschlagener Versuch, Selbstbedienung im Supermarkt) sind unabhängig von ihrer Punktzahl im Plan.")
a(f"   - Grundlagen kommen vor den Themen, die auf ihnen aufbauen: {ST['hart']} verbindliche Abhängigkeiten, {ST['hart_verletzt']} Verstöße. Beispiele: Verwaltungsakt → Klagearten → Anfechtungsklage → § 80 V → Nachbareilverfahren; Straftataufbau → Notwehr; Versuch → Rücktritt.")
a(f"   - Fallvideos (Montag) sind davon ausgenommen, weil sie ihren Hintergrund selbst erklären. In {ST['weich_vorher']} Fällen läuft ein Fallvideo vor dem passenden Schemavideo; dort verweist die Folge am Ende auf das Schema (Karte oder Link), sobald es erschienen ist.")
a("5. **Verteilung auf die Jahre:**")
a("   - Die stärksten Themen stehen vorn, jedes Jahr hat dieselbe Gebietsmischung.")
a("   - Jahr 1 und 2 enthalten die meisten Klassiker; die späteren Jahre vertiefen (Irrtümer im Detail, Nebengebiete, Sonderlagen des 2. Examens).")
a("   - In den Klausurphasen (Januar/Februar, Juni/Juli) bevorzugt der Freitag Abgrenzungen und Klausurfehler.\n")
a("| Jahr | Ø Punkte | Klassiker | Themen mit hoher Breite (≥ 4) | Charakter |")
a("| ---: | ---: | ---: | ---: | --- |")
CH = {1: "Fundament und Zugpferde", 2: "Klassiker-Ausbau", 3: "Vertiefung der Kerngebiete", 4: "Streitstände und Nebengebiete", 5: "Spezialprobleme, Feinschliff"}
for j, s, k, b in jahrstat:
    a(f"| {j} | {s} | {k} | {b} | {CH[j]} |")
a("")
a("## Wachstumshebel (aus der Recherche)\n")
a("- **Shorts als Zubringer:** Aus jeder Folge 1–3 Shorts schneiden, etwa Merksatz, Fallfrage oder die überraschende Wendung. Laut Metricool-Studie wuchsen Shorts-Aufrufe bei Kanälen unter 10.000 Abonnenten um rund 200 %; 780 Folgen ergeben so 1.500–2.000 Shorts ohne neue Inhalte.")
a("- **Playlists je Reihe und Gebiet:** Zum Beispiel „Der Fall – Strafrecht“, „Prüfungsschemata Zivilrecht“ oder „2. Examen: Revision“. Die Grundstock-Playlist dient als Einsteigerpfad fürs erste Semester.")
a("- **Titel als Frage mit konkretem Bild:** Etwa „Darf …?“, „Muss …?“ oder „Mord mit dem Auto?“; dazu der Normbezug im Untertitel oder in der Beschreibung für die Suche.")
a("- **Aktualitäts-Joker:** Die Vorproduktion schließt Aktualität nicht aus. Bei einem großen neuen BGH- oder BVerfG-Urteil, einer Reform oder der BfJ-Examensstatistik (jährlich Ende Juni) ersetzt ein kurzfristiges Video einen Freitagsslot. Das verdrängte Thema rückt an das Ende des Jahres. Aktualitätsthemen brachten WBS zeitweise rund 20.000 Abonnenten pro Woche.")
a("- **Saisonalität:** In den drei Wochen vor Prüfungen steigt die Lernaktivität um 56–74 % (Jurafuchs). Schemata und Klausurfehler deshalb im Januar/Februar und Juni/Juli, Grundlagen und Methodik im Oktober und April.")
a("- **Zielgruppe:**")
a("  - WS 2024/25: 115.272 Jurastudierende (Destatis).")
a("  - 2024 haben 9.255 die erste und 8.084 die zweite Prüfung bestanden (BfJ).")
a("  - Der Kanal zielt auf diese Gruppe. Brückenthemen mit Breite 5 (Polizei, Kauf, Miete, Erbe, Verkehr) öffnen ihn für Laien.\n")
a("## Vor der Produktion: Qualitäts- und Rechtsstandsregeln\n")
a(f"- **Rechtsstand:** {len(rs)} Folgen tragen einen Rechtsstandshinweis, etwa MoPeG 2024, Kaufrechtsreform, Recht auf Reparatur seit 23.07.2026, Cannabisgesetz, Vier-Tages-Fiktion, AG-Zuständigkeit bis 10.000 €, § 246e BauGB oder die Verortung der BVerfG-Zuständigkeiten nach der GG-Änderung 2024 (laut Repetitorium Art. 94 GG, am amtlichen Text zu prüfen). "
  f"Diese Folgen stehen überwiegend in den späteren Jahren. Sie werden erst **kurz vor der Veröffentlichung** geskriptet oder vor dem Upload erneut geprüft. Stabile Themen zuerst vorproduzieren.")
a(f"- **Länderrecht:** {len(laender)} Folgen betreffen Polizei-, Versammlungs-, Bau- oder Kommunalrecht. Sie werden bundesweit erklärt; die Landesnormen stehen als Tabelle auf der Karte.")
a("- **Fundstellen:** Die Angaben unter „Leitentscheidung“ (Aktenzeichen, BGHSt-/BVerfGE-Fundstellen) stammen teils nicht aus dem Repetitorium. Sie müssen vor dem Skript an der Primärquelle geprüft werden, das gilt vor allem für die in den Auswertungen als unsicher markierten.")
a("- **Fehler im Repetitorium, bei der Auswertung gefunden:**")
a("  - Teil C Fall 49 ordnet den Hoferben-Fall (BGHSt 37, 214) als aberratio ictus ein. Tatsächlich geht es um einen error in persona des Täters und seine Folgen für den Anstifter.")
a("  - Band 14 Nr. 13 beschreibt § 339 StPO falsch; gemeint ist § 302 I 2 StPO. Band 6 Nr. 46 ist richtig.")
a("- **Eigenständigkeit:** Die Einstiege sind eigene Fälle. Die Fallnamen und Sachverhalte der Lernplattform, deren Struktur das Repetitorium spiegelt, werden nicht verwendet.")
a("- Jede Folge durchläuft weiterhin den [Produktionsstandard](MASTERSTANDARD-09.md) und den [Abnahmebogen](ABNAHME-16x9.md) mit Rechtscheck am amtlichen Normtext.\n")
a("## Jahr 1 im Überblick: die ersten zwölf Wochen\n")
a("| Woche | Montag · Der Fall | Mittwoch · Examenswissen | Freitag · Klausurpraxis |")
a("| ---: | --- | --- | --- |")
for w in range(1, 13):
    z = {x["tag"]: x for x in P if x["jahr"] == 1 and x["woche_im_jahr"] == w}
    a(f"| {w} | " + " | ".join(md(z[t]["titel"]) + f" <sub>{z[t]['gebiet']}</sub>" for t in ("Mo", "Mi", "Fr")) + " |")
a("")
a("## Alle 780 Folgen\n")
a("Spalten: Nr., Woche (fortlaufend), Tag, Gebiet, Titel, zentrale Normen, Leitentscheidung. Einstieg, Kernfrage, Bewertung, Voraussetzung und Fundstelle stehen in der CSV.\n")
for j in range(1, 6):
    a(f"### Jahr {j}\n")
    a("| Nr. | Wo. | Tag | Gebiet | Titel | Normen | Leitentscheidung |")
    a("| ---: | ---: | --- | --- | --- | --- | --- |")
    for x in P:
        if x["jahr"] == j:
            a(f"| {x['nr']} | {x['woche']} | {x['tag']} | {x['gebiet']} | {md(x['titel'])} | {md(x['normen'])} | {md(x['leitentscheidung'])} |")
    a("")
open(ZIEL_MD, "w").write("\n".join(L) + "\n")
print("geschrieben:", ZIEL_MD, len(P), "Folgen,", len(R), "Reserve")
